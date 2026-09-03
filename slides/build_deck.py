"""Genera un PPTX instructor-oriented (sin diseño) para guiar el setup de
Microsoft Foundry y los labs de Anders (Lab 3) y Julie (Lab 4)."""

from pptx import Presentation
from pptx.util import Pt

prs = Presentation()  # plantilla por defecto (16:9? no; por defecto 4:3). Ajustamos a 16:9.
prs.slide_width = Pt(960)
prs.slide_height = Pt(540)

TITLE_SLIDE = prs.slide_layouts[0]
TITLE_AND_CONTENT = prs.slide_layouts[1]
SECTION_HEADER = prs.slide_layouts[2]


def add_title(title, subtitle=""):
    slide = prs.slides.add_slide(TITLE_SLIDE)
    slide.shapes.title.text = title
    if subtitle:
        slide.placeholders[1].text = subtitle
    return slide


def add_section(title):
    slide = prs.slides.add_slide(SECTION_HEADER)
    slide.shapes.title.text = title
    return slide


def add_bullets(title, bullets, notes=""):
    """bullets: lista de (texto, nivel)."""
    slide = prs.slides.add_slide(TITLE_AND_CONTENT)
    slide.shapes.title.text = title
    body = slide.placeholders[1].text_frame
    body.clear()
    for i, item in enumerate(bullets):
        text, level = item if isinstance(item, tuple) else (item, 0)
        p = body.paragraphs[0] if i == 0 else body.add_paragraph()
        p.text = text
        p.level = level
    if notes:
        slide.notes_slide.notes_text_frame.text = notes
    return slide


# ---------------------------------------------------------------------------
# 1. Portada
# ---------------------------------------------------------------------------
add_title(
    "Taller Multi-Agéntico — Guía del Instructor",
    "Setup de Microsoft Foundry · Lab 3 (Anders) · Lab 4 (Julie)",
)

# ---------------------------------------------------------------------------
# 2. Objetivo de la sesión
# ---------------------------------------------------------------------------
add_bullets(
    "Objetivo de esta sesión",
    [
        ("Guiar a los asistentes por la capa de razonamiento del taller (Foundry).", 0),
        ("Desplegar la infraestructura base en la suscripción de cada participante.", 0),
        ("Construir y ejecutar dos agentes con roles complementarios:", 0),
        ("Anders — agente ejecutor (renderiza reportes).", 1),
        ("Julie — agente planner/orquestador (campañas de marketing).", 1),
        ("Dejar a Anders listo para consumirse desde Copilot Studio (Lab 8).", 0),
    ],
    notes=(
        "Enmarca la sesión: esta es la capa del medio (Foundry). Fabric ya entregó "
        "los datos; Copilot Studio será la capa de interacción. Recalca que Anders y "
        "Julie son COMPLEMENTARIOS: uno ejecuta/renderiza, el otro planifica/consulta. "
        "Duración sugerida: ~15 min setup, ~25 min Anders, ~30 min Julie."
    ),
)

# ---------------------------------------------------------------------------
# 3. Arquitectura de 3 capas
# ---------------------------------------------------------------------------
add_bullets(
    "Arquitectura de tres capas",
    [
        ("Copilot Studio  →  capa de interacción (Bill, Charlie, Ric).", 0),
        ("Microsoft Foundry  →  capa de razonamiento (Anders, Julie).  ← HOY", 0),
        ("Microsoft Fabric  →  capa de datos (Mark, Amy).", 0),
        ("Modelos GPT-5.1 desplegados en Azure AI Services.", 0),
        ("Punto de unión: la Function App `FxContosoRetail` expone la API vía OpenAPI.", 0),
    ],
    notes=(
        "Dibuja mentalmente el sándwich de 3 capas. Hoy trabajamos la del medio. "
        "Insiste en que la Function App es el puente: Anders y Julie no hablan con SQL "
        "ni con Fabric directamente; consumen la API por su especificación OpenAPI. "
        "Mark (Fabric) obtiene datos; Bill (Copilot) orquesta; Anders/Julie razonan."
    ),
)

# ---------------------------------------------------------------------------
# 4. Anders vs Julie (mapa mental)
# ---------------------------------------------------------------------------
add_bullets(
    "Los dos agentes de hoy",
    [
        ("ANDERS — Executor Agent", 0),
        ("kind: \"prompt\" + 1 herramienta OpenAPI.", 1),
        ("Recibe datos ya listos y RENDERIZA el reporte (endpoint ordersReporter).", 1),
        ("NO consulta la base de datos.", 1),
        ("Se conecta a Copilot Studio (necesita Activity Protocol).", 1),
        ("JULIE — Planner Workflow", 0),
        ("kind: \"workflow\" + 3 herramientas (2 sub-agentes + 1 OpenAPI).", 1),
        ("CONSULTA la base (genera T-SQL y lo ejecuta vía SqlExecutor).", 1),
        ("Orquestador AUTOSUFICIENTE dentro de Foundry (no usa Copilot).", 1),
    ],
    notes=(
        "Esta es la diapositiva ancla de toda la sesión. Regla mnemotécnica: "
        "ANDERS RENDERIZA (recibe datos, arma reporte, sin BD) · JULIE CONSULTA "
        "(genera y ejecuta SQL, con BD). Por eso solo Julie necesita permisos de BD, "
        "y solo Anders necesita el Activity Protocol (porque es el que Copilot invoca)."
    ),
)

# ---------------------------------------------------------------------------
# 5. SECCIÓN: SETUP
# ---------------------------------------------------------------------------
add_section("Parte 0 — Setup de Foundry")

# ---------------------------------------------------------------------------
# 6. Setup: opción recomendada + prerrequisitos
# ---------------------------------------------------------------------------
add_bullets(
    "Setup — Punto de partida",
    [
        ("Recomendado: GitHub Codespaces (entorno listo en ~2 min, sin instalar nada).", 0),
        ("Alternativa local, prerrequisitos:", 0),
        ("Azure CLI actualizado.", 1),
        (".NET 8 SDK.", 1),
        ("PowerShell 7+ (pwsh) — NO PowerShell 5.1.", 1),
        ("Suscripción Azure (Owner/Contributor) del tenant temporal asignado.", 1),
        ("Anotar el número de tenant (ej. usuario@azurehol3387.com → 3387).", 0),
    ],
    notes=(
        "Empuja Codespaces salvo que alguien insista en local. GOTCHA #1: los scripts "
        "SON de PowerShell 7 (pwsh); en 5.1 fallan. En la terminal por defecto de "
        "Codespaces (bash) hay que invocar los .ps1 con 'pwsh ./script.ps1'. "
        "El número de tenant genera el suffix único de los recursos."
    ),
)

# ---------------------------------------------------------------------------
# 7. Setup: despliegue
# ---------------------------------------------------------------------------
add_bullets(
    "Setup — Despliegue de infraestructura",
    [
        ("Automatizado con Bicep + PowerShell (sin tocar el portal).", 0),
        ("1) az login  →  2) az account show  →  confirmar suscripción correcta.", 0),
        ("3) az bicep upgrade", 0),
        ("4) cd labs/foundry/setup/op-flex  →  ./deploy.ps1", 0),
        ("El script pregunta: TenantName, Location, ResourceGroup, y (opcional) SQL de Fabric.", 0),
        ("En < 10 min queda el ambiente completo.", 0),
    ],
    notes=(
        "Demuestra el login primero: si el usuario elige mal la suscripción, todo el "
        "resto falla. El script es interactivo: Enter acepta defaults (eastus, "
        "rg-contoso-retail). Si preguntan por la SQL de Fabric y aún no la tienen, "
        "pueden decir 'N' y configurarla luego (solo la necesita Julie)."
    ),
)

# ---------------------------------------------------------------------------
# 8. Setup: recursos + suffix
# ---------------------------------------------------------------------------
add_bullets(
    "Setup — Recursos creados y el suffix",
    [
        ("stcontosoretail{suffix}  — Storage de la Function App.", 0),
        ("asp-contosoretail-{suffix}  — App Service Plan (Flex).", 0),
        ("func-contosoretail-{suffix}  — Function App (API, .NET 8 isolated).", 0),
        ("ais-contosoretail-{suffix}  — AI Foundry Resource (GPT-5.1).", 0),
        ("aip-contosoretail-{suffix}  — AI Foundry Project.", 0),
        ("Sacar el suffix rápido:", 0),
        ("az cognitiveservices account list -g rg-contoso-retail --query \"[0].name\" -o tsv", 1),
    ],
    notes=(
        "El {suffix} son 5 caracteres derivados del número de tenant; evita colisiones "
        "entre participantes. Aparece en TODOS los nombres y lo usarán en varios pasos "
        "(appsettings, permisos SQL, publicación de Anders). Enséñales a extraerlo con "
        "el comando de la última línea o tomando el último segmento del nombre del recurso."
    ),
)

# ---------------------------------------------------------------------------
# 9. Setup: RBAC + gotchas
# ---------------------------------------------------------------------------
add_bullets(
    "Setup — RBAC y advertencias del instructor",
    [
        ("RBAC obligatorio: rol \"Cognitive Services User\" sobre ais-contosoretail-{suffix}.", 0),
        ("Sin él: error PermissionDenied al crear/ejecutar agentes.", 1),
        ("Propagación RBAC: hasta ~1 min de espera.", 0),
        ("Gotchas frecuentes:", 0),
        ("Ejecutar .ps1 con pwsh (en bash: 'pwsh ./script.ps1').", 1),
        ("Windows: Set-ExecutionPolicy RemoteSigned -Scope CurrentUser (una vez).", 1),
        ("Storage bloqueado por política (error 503) → ver unlock-storage.ps1.", 1),
    ],
    notes=(
        "El RBAC es la causa #1 de soporte en este punto: si no asignan Cognitive "
        "Services User a SU usuario, no pueden crear agentes. Recuérdales esperar el "
        "minuto de propagación antes de correr los labs. Menciona unlock-storage.ps1 "
        "por si alguien topa con el 503 del storage por política del tenant."
    ),
)

# ---------------------------------------------------------------------------
# 10. SECCIÓN: LAB 3 ANDERS
# ---------------------------------------------------------------------------
add_section("Lab 3 — Anders (Executor Agent)")

# ---------------------------------------------------------------------------
# 11. Anders: concepto
# ---------------------------------------------------------------------------
add_bullets(
    "Lab 3 — Concepto de Anders",
    [
        ("Rol: recibir acciones operativas y EJECUTARLAS (renderizar reportes de órdenes).", 0),
        ("Herramienta: OpenAPI Tool generada desde la spec de la Function App.", 0),
        ("El agente descubre e invoca los endpoints automáticamente.", 1),
        ("Endpoint que usa: ordersReporter (arma y devuelve la URL del reporte).", 0),
        ("Anders NO toca SQL: recibe la lista de órdenes ya construida.", 0),
    ],
    notes=(
        "Deja clarísimo el límite de responsabilidad: Anders es un 'ejecutor de "
        "acciones', no un consultor de datos. En el flujo de Copilot, Mark obtiene las "
        "órdenes y Bill se las pasa a Anders; Anders solo las transforma al schema del "
        "endpoint y las renderiza. Esto explica por qué no necesita permisos de BD."
    ),
)

# ---------------------------------------------------------------------------
# 12. Anders: flujo del código
# ---------------------------------------------------------------------------
add_bullets(
    "Lab 3 — Cómo funciona el código",
    [
        ("Fase 1: descarga la spec OpenAPI de la Function App ({base}/openapi/v3.json).", 0),
        ("Fase 2: crea (o reutiliza) el agente en Foundry con la OpenAPI Tool.", 0),
        ("Instrucciones: construir el JSON con el schema EXACTO de ordersReporter.", 1),
        ("Fase 3: chat interactivo con la Responses API.", 0),
        ("Dos versiones de SDK en /code: ms-foundry (recomendada) y ai-foundry.", 0),
    ],
    notes=(
        "Muestra el Program.cs: la clave pedagógica es que el agente se auto-configura "
        "leyendo la spec OpenAPI — no se hardcodean endpoints. Señala el prompt de "
        "Anders con el schema estricto (orderNumber, orderLineNumber secuencial, fechas "
        "ISO, numéricos). Recomienda la carpeta ms-foundry."
    ),
)

# ---------------------------------------------------------------------------
# 13. Anders: pasos guiados
# ---------------------------------------------------------------------------
add_bullets(
    "Lab 3 — Pasos guiados",
    [
        ("3.1  Verificar soporte OpenAPI en FxContosoRetail (ya viene preconfigurado).", 0),
        ("Paquetes: OpenApi + Microsoft.Data.SqlClient · Endpoints: OrdersReporter, SqlExecutor.", 1),
        ("3.2  Verificar la spec OpenAPI (JSON / Swagger UI).", 0),
        ("3.3  Configurar appsettings.json, compilar, ejecutar y probar Anders.", 0),
        ("appsettings: FoundryProjectEndpoint, ModelDeploymentName, FunctionAppBaseUrl.", 1),
    ],
    notes=(
        "En 3.1 aclara que SqlClient y SqlExecutor pertenecen a la Function App (los usa "
        "Julie), NO a Anders. Anders solo usa OrdersReporter. En 3.3 el error típico es "
        "un appsettings mal copiado: el endpoint del proyecto lleva /api/projects/aip-... "
        "Prueba con un caso real y muestra la URL del reporte generado."
    ),
)

# ---------------------------------------------------------------------------
# 14. Anders: publicar para Copilot (Activity Protocol)
# ---------------------------------------------------------------------------
add_bullets(
    "Lab 3 — Publicar Anders para Copilot Studio (Activity Protocol)",
    [
        ("Copilot Studio invoca agentes de Foundry con el Activity Protocol.", 0),
        ("Por defecto un agente solo expone Responses API → falla con HTTP 400.", 0),
        ("\"endpoint does not support activity\".", 1),
        ("Opción A (portal): Publish → Teams y Microsoft 365 → Direct publish → Just you.", 0),
        ("Opción B (sin portal, automatizada):", 0),
        ("pwsh ./setup/enable-activity-protocol.ps1 -Suffix <suffix>", 1),
        ("Habilita 'activity' conservando 'responses' + esquemas Entra y BotServiceRbac.", 0),
    ],
    notes=(
        "PUNTO CRÍTICO y relativamente nuevo: Microsoft cambió el modelo de publicación "
        "de Foundry, por eso 'antes no salía'. Si no se habilita Activity, la conexión "
        "de Bill (Lab 8) falla con 400. El script hace el PATCH REST (api-version=v1, "
        "preview) sin necesidad del portal. Verificable con un GET al agente: debe "
        "listar protocols [activity, responses]. Solo Anders lo necesita, NO Julie."
    ),
)

# ---------------------------------------------------------------------------
# 15. SECCIÓN: LAB 4 JULIE
# ---------------------------------------------------------------------------
add_section("Lab 4 — Julie (Planner Workflow)")

# ---------------------------------------------------------------------------
# 16. Julie: concepto
# ---------------------------------------------------------------------------
add_bullets(
    "Lab 4 — Concepto de Julie",
    [
        ("Rol: orquestar campañas de marketing personalizadas.", 0),
        ("Tipo: agente workflow (orquestador AUTOSUFICIENTE, sin Copilot).", 0),
        ("Coordina 3 herramientas: SqlAgent, MarketingAgent y la OpenAPI (SqlExecutor).", 0),
        ("Julie SÍ consulta la base: genera T-SQL y lo ejecuta.", 0),
    ],
    notes=(
        "Contrasta con Anders: Julie fue diseñada como orquestador que vive y decide "
        "dentro de Foundry, invocando sub-agentes por código (SDK). Nunca la llama "
        "Copilot Studio, por eso NO necesita Activity Protocol. Pero sí toca datos, "
        "por eso SÍ necesita permisos de BD (siguiente lámina)."
    ),
)

# ---------------------------------------------------------------------------
# 17. Julie: flujo de 5 pasos
# ---------------------------------------------------------------------------
add_bullets(
    "Lab 4 — El workflow de Julie (5 pasos)",
    [
        ("1) Extrae el filtro/segmento de clientes de la petición en lenguaje natural.", 0),
        ("2) Invoca a SqlAgent → genera el T-SQL.", 0),
        ("3) Ejecuta la consulta contra Fabric vía Function App (SqlExecutor / OpenAPI).", 0),
        ("4) Invoca a MarketingAgent (con Bing Search) → mensaje por cliente.", 0),
        ("5) Organiza la salida como JSON de campaña de correos.", 0),
    ],
    notes=(
        "Recorre el flujo de punta a punta: entra un segmento ('clientes que compraron X "
        "el último mes'), sale un JSON de campaña con mensajes personalizados. El paso 3 "
        "es el que requiere permisos de BD. El paso 4 usa Bing, útil para mensajes con "
        "contexto actual (tema del Challenge 1)."
    ),
)

# ---------------------------------------------------------------------------
# 18. Julie: sub-agentes
# ---------------------------------------------------------------------------
add_bullets(
    "Lab 4 — Sub-agentes de Julie",
    [
        ("SqlAgent.cs — traduce lenguaje natural a T-SQL y ejecuta vía SqlExecutor.", 0),
        ("MarketingAgent.cs — redacta mensajes personalizados apoyado en Bing.", 0),
        ("JulieAgent.cs — define a Julie como workflow (CSDL YAML) e invoca sub-agentes.", 0),
        ("Program.cs — carga config, crea/verifica agentes en Foundry y ejecuta el chat.", 0),
    ],
    notes=(
        "Abre la carpeta code/agents/JulieAgent. La orquestación es de tipo 'workflow' "
        "(definida en CSDL YAML) en contraste con un agente 'prompt' simple: permite "
        "coordinar sub-agentes y tools en un flujo determinístico. Muéstralo como "
        "patrón reutilizable de multi-agente dentro de Foundry."
    ),
)

# ---------------------------------------------------------------------------
# 19. Julie: permisos de BD (CRÍTICO)
# ---------------------------------------------------------------------------
add_bullets(
    "Lab 4 — Permisos de BD (obligatorio, causa #1 de fallos)",
    [
        ("El acceso a datos va: Julie → Function App (Managed Identity) → SQL de Fabric.", 0),
        ("Parte A — Acceso al workspace de Fabric:", 0),
        ("Manage access → Add → func-contosoretail-{suffix} → rol Contributor.", 1),
        ("Parte B — Usuario SQL en la base retail (New Query):", 0),
        ("CREATE USER [func-contosoretail-{suffix}] FROM EXTERNAL PROVIDER;", 1),
        ("ALTER ROLE db_datareader ADD MEMBER [func-contosoretail-{suffix}];", 1),
        ("Esperar 1–3 min de propagación antes de probar.", 0),
    ],
    notes=(
        "Este es el paso que más soporte genera en Julie. La Function App se autentica "
        "con su Managed Identity (auth 'Active Directory Default', sin secretos), así que "
        "hay que autorizar ESA identidad tanto en el workspace (Parte A) como dentro de "
        "la base con un usuario externo + db_datareader (Parte B). Si preguntan por qué "
        "solo lectura: Julie solo consulta, no escribe. Recuerda los 1–3 min de espera."
    ),
)

# ---------------------------------------------------------------------------
# 20. Julie: pasos guiados
# ---------------------------------------------------------------------------
add_bullets(
    "Lab 4 — Pasos guiados",
    [
        ("Paso 1: configurar appsettings.json (endpoint, modelo, base URL, SQL de Fabric).", 0),
        ("Paso 2: verificar que los permisos de Fabric (Partes A y B) están listos.", 0),
        ("Paso 3: ejecutar Julie.", 0),
        ("Paso 4: probar el flujo end-to-end → JSON de campaña.", 0),
        ("Challenges: mejorar el prompt de MarketingAgent · agente no-code con Code Interpreter.", 0),
    ],
    notes=(
        "Los valores SQL (FabricWarehouseSqlEndpoint, FabricWarehouseDatabase) salen del "
        "connection string ADO.NET del Warehouse (guía sql-parameters.md): endpoint = "
        "Data Source sin ',1433'; database = Initial Catalog. Si alguien no desplegó "
        "Fabric, puede apuntar a un Azure SQL standalone. Cierra con los challenges si "
        "hay tiempo."
    ),
)

# ---------------------------------------------------------------------------
# 21. Resumen mental
# ---------------------------------------------------------------------------
add_bullets(
    "Resumen para no confundirse",
    [
        ("ANDERS = RENDERIZA. Recibe datos, arma reporte (ordersReporter). Sin BD.", 0),
        ("Anders → sí necesita Activity Protocol (lo invoca Copilot/Bill).", 1),
        ("JULIE = CONSULTA. Genera y ejecuta SQL (SqlExecutor). Con BD.", 0),
        ("Julie → orquestador autosuficiente; NO necesita Activity Protocol.", 1),
        ("Solo los agentes de Foundry conectados en Copilot como 'agente externo' usan Activity.", 0),
    ],
    notes=(
        "Cierra la sesión repitiendo la regla ancla. Es la pregunta que más surge: "
        "¿por qué solo Anders necesita el script de Activity? Porque es el único que "
        "Copilot Studio invoca externamente. Julie se queda en Foundry."
    ),
)

# ---------------------------------------------------------------------------
# 22. Checklist del instructor
# ---------------------------------------------------------------------------
add_bullets(
    "Checklist del instructor",
    [
        ("Setup: az login correcto · deploy.ps1 OK · RBAC Cognitive Services User asignado.", 0),
        ("Suffix identificado y a la mano.", 0),
        ("Anders: appsettings OK · corre y genera reporte · Activity Protocol habilitado.", 0),
        ("Julie: permisos Fabric (A y B) · SQL params en appsettings · flujo end-to-end OK.", 0),
        ("Recordatorio: ejecutar .ps1 con pwsh; esperar propagaciones RBAC/SQL.", 0),
    ],
    notes=(
        "Úsala como lista de verificación en vivo. Si algo falla, el 90% de las veces es: "
        "(1) suscripción equivocada en az login, (2) RBAC no asignado o sin propagar, "
        "(3) permisos de BD faltantes para Julie, (4) ejecutar el .ps1 con bash en vez "
        "de pwsh."
    ),
)

out = "slides/Guia-Instructor-Foundry-Anders-Julie.pptx"
prs.save(out)
print(f"OK: {out} — {len(prs.slides.__iter__.__self__._sldIdLst)} slides")
